"""SpringerOpen (Geothermal Energy) manuscript helper.

Times New Roman 12 pt, double spacing, continuous line numbers, page numbers, no page breaks, tables without
shading, figure/table titles + legends in text, Vancouver-numbered citations [@key].
Inline markup in text: *italic*, **bold**, x_{sub}, x^{sup} (plain Unicode runs, robust to PDF conversion),
$latex$ for an inline equation object (used sparingly). Display equations via eq() are Word (OMML) equation
objects, which remain editable.
"""
import re
from lxml import etree
from latex2mathml.converter import convert as l2m
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

XSL = etree.XSLT(etree.parse(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"))
FONT = "Times New Roman"


def omml(latex):
    return XSL(etree.fromstring(l2m(latex).encode())).getroot()


class SDoc:
    def __init__(self, line_numbers=True, spacing=2.0, size=12):
        self.d = Document()
        s = self.d.sections[0]
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.5)
        s.top_margin = s.bottom_margin = Cm(2.5)
        st = self.d.styles["Normal"]
        st.font.name = FONT; st.font.size = Pt(size)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        pf = st.paragraph_format
        pf.space_after = Pt(0); pf.space_before = Pt(0); pf.line_spacing = spacing
        self.size = size
        for lvl, sz in ((1, size + 2), (2, size), (3, size)):
            h = self.d.styles[f"Heading {lvl}"]
            h.font.name = FONT; h.font.size = Pt(sz); h.font.bold = True; h.font.italic = (lvl == 3)
            h.font.color.rgb = None
            h.paragraph_format.space_before = Pt(12); h.paragraph_format.space_after = Pt(0)
            h.paragraph_format.line_spacing = spacing
            rpr = h.element.get_or_add_rPr()
            rf = rpr.find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts"); rpr.append(rf)
            for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(a), FONT)
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rf.get(qn(a)) is not None:
                    del rf.attrib[qn(a)]
            for tag in ("w:color",):
                e = rpr.find(qn(tag))
                if e is not None:
                    rpr.remove(e)
        self.fig_n = self.tab_n = self.eq_n = 0
        self.cite_order = []
        p = s.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._field(p, "PAGE")
        if line_numbers:
            ln = OxmlElement("w:lnNumType")
            ln.set(qn("w:countBy"), "1"); ln.set(qn("w:restart"), "continuous"); ln.set(qn("w:distance"), "300")
            s._sectPr.append(ln)

    def _field(self, p, code):
        r = p.add_run()
        for t, v in (("begin", None), ("instr", code), ("end", None)):
            if t == "instr":
                e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = v
            else:
                e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t)
            r._r.append(e)

    # ------------------------------------------------------------------ citations
    def cite(self, text):
        def rep(m):
            keys = [k.strip().lstrip("@") for k in m.group(1).split(";")]
            nums = []
            for k in keys:
                if k not in self.cite_order:
                    self.cite_order.append(k)
                nums.append(self.cite_order.index(k) + 1)
            nums = sorted(set(nums)); out, i = [], 0
            while i < len(nums):
                j = i
                while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
                    j += 1
                out.append(f"{nums[i]}–{nums[j]}" if j - i >= 2 else ",".join(str(n) for n in nums[i:j + 1]))
                i = j + 1
            return "[" + ",".join(out) + "]"
        return re.sub(r"\[(@[^\]]+)\]", rep, text)

    # ------------------------------------------------------------------ text
    TOK = re.compile(r"(\$[^$]+\$|\*\*[^*]+\*\*|(?<![A-Za-z0-9])\*[^*\s][^*]*\*|_\{[^}]+\}|\^\{[^}]+\})")

    # hyphen used as a sign before a number -> minus sign (not after a word character, ')' or en dash,
    # so "56-32", "16A(78)-32" and "5-m" are left alone)
    MINUS = re.compile(r"(?<![\w\)–])-(?=\d)")

    def runs(self, p, text, size=None, bold=False, italic=False):
        text = self.MINUS.sub("−", self.cite(text))
        for tk in self.TOK.split(text):
            if not tk:
                continue
            if tk.startswith("$") and tk.endswith("$") and len(tk) > 2:
                p._p.append(omml(tk[1:-1])); continue
            b, i, sub, sup = bold, italic, False, False
            if tk.startswith("**") and tk.endswith("**") and len(tk) > 4:
                tk, b = tk[2:-2], True
            elif tk.startswith("*") and tk.endswith("*") and len(tk) > 2:
                tk, i = tk[1:-1], True
            elif tk.startswith("_{"):
                tk, sub = tk[2:-1], True
            elif tk.startswith("^{"):
                tk, sup = tk[2:-1], True
            r = p.add_run(tk)
            r.bold, r.italic = b, i
            r.font.subscript, r.font.superscript = sub, sup
            if size:
                r.font.size = Pt(size)
        return p

    def p(self, text, align="justify", size=None, keep=False, **kw):
        p = self.d.add_paragraph()
        p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "left": WD_ALIGN_PARAGRAPH.LEFT,
                       "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
        if keep:
            p.paragraph_format.keep_with_next = True
        self.runs(p, text, size=size, **kw)
        return p

    def h(self, text, level=1):
        return self.d.add_heading(text, level=level)

    def bullets(self, items):
        for it in items:
            p = self.d.add_paragraph(style="List Bullet")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            self.runs(p, it)

    def eq(self, latex):
        """Display equation (OMML oMathPara) in a borderless two-column table with the number on the same row."""
        self.eq_n += 1
        sec = self.d.sections[0]
        width = (sec.page_width - sec.left_margin - sec.right_margin) / 360000
        t = self.d.add_table(rows=1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
        c0, c1 = t.rows[0].cells
        c0.width, c1.width = Cm(width - 1.6), Cm(1.6)
        p = c0.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para = OxmlElement("m:oMathPara"); para.append(omml(latex)); p._p.append(para)
        q = c1.paragraphs[0]; q.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        q.add_run(f"({self.eq_n})")
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        for c in (c0, c1):
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        return self.eq_n

    embed = False

    def figure(self, title, legend, file=None):
        """SpringerOpen: figures are uploaded separately; title (<= 15 words) and legend go in the text.
        With embed = True (review copy) the figure file is inserted above its caption."""
        self.fig_n += 1
        assert len(title.split()) <= 15, title
        if self.embed and file is not None:
            self.picture(file, 16.5)
        p = self.d.add_paragraph()
        r = p.add_run(f"Fig. {self.fig_n} "); r.bold = True
        self.runs(p, title, bold=True)
        self.p(legend)
        return self.fig_n

    def table(self, title, header, rows, widths_cm, legend=None, size=10, label=None):
        if label is None:
            self.tab_n += 1; label = f"Table {self.tab_n}"
        assert len(title.split()) <= 15, title
        c = self.d.add_paragraph(); c.paragraph_format.keep_with_next = True
        r = c.add_run(f"{label} "); r.bold = True
        self.runs(c, title, bold=True)
        t = self.d.add_table(rows=1 + len(rows), cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
        self._borders(t)
        for j, htxt in enumerate(header):
            cell = t.rows[0].cells[j]; cell.width = Cm(widths_cm[j])
            pp = cell.paragraphs[0]; pp.paragraph_format.line_spacing = 1.0
            self.runs(pp, htxt, size=size, bold=True)
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                cell = t.rows[i + 1].cells[j]; cell.width = Cm(widths_cm[j])
                pp = cell.paragraphs[0]; pp.paragraph_format.line_spacing = 1.0
                self.runs(pp, str(v), size=size)
        if legend:
            self.p(legend, size=size)
        return label

    def _borders(self, t):
        tblPr = t._tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for edge, sz in (("top", 8), ("bottom", 8), ("insideH", 4)):
            e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:color"), "000000")
            b.append(e)
        for edge in ("left", "right", "insideV"):
            e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "nil"); b.append(e)
        tblPr.append(b)

    def picture(self, path, width_cm=16):
        p = self.d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Cm(width_cm))

    def save(self, path):
        self.d.save(str(path)); print("saved", path)
