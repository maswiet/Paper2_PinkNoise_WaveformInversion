"""Geophysical Journal International (GJI) manuscript helper, built on springer_doc.SDoc.

Differences from SDoc:
* author-year citations: [@a; @b] -> (Author 2001; Author & Other 2002); {@a} -> Author et al. (2001)
  (narrative); <@a> -> Author et al. 2001 (bare, inside existing parentheses); the reference list is alphabetical and uses the GJI reference style;
* numbered GJI headings: level 1 upper case ("1 INTRODUCTION"), level 2 "1.1 Title", level 3 italic;
* display equations are ordinary paragraphs (centred equation, right-aligned number), not tables;
* captions "Figure n." and "Table n."
REFS passed to GDoc(refs=...) maps key -> (in-text label, full reference), e.g.
  "aki1975": ("Aki & Chouet 1975", "Aki, K. & Chouet, B., 1975. Origin of coda waves ..., J. geophys. Res., 80, 3322-3342.")
"""
import re
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from springer_doc import SDoc, omml


class GDoc(SDoc):
    def __init__(self, refs, **kw):
        super().__init__(**kw)
        self.refs = refs
        self.used = []
        self.hnum = [0, 0, 0]

    # ------------------------------------------------------------------ citations
    def _lab(self, k):
        k = k.strip().lstrip("@")
        if k not in self.refs:
            raise KeyError(f"unknown reference {k}")
        if k not in self.used:
            self.used.append(k)
        return self.refs[k][0]

    def cite(self, text):
        def par(m):
            return "(" + "; ".join(self._lab(k) for k in m.group(1).split(";")) + ")"

        def nar(m):
            out = []
            for k in m.group(1).split(";"):
                lab = self._lab(k)
                name, year = lab.rsplit(" ", 1)
                out.append(f"{name} ({year})")
            return "; ".join(out)
        text = re.sub(r"\[(@[^\]]+)\]", par, text)
        text = re.sub(r"<(@[^>]+)>", lambda m: "; ".join(self._lab(k) for k in m.group(1).split(";")), text)
        return re.sub(r"\{(@[^}]+)\}", nar, text)

    # ------------------------------------------------------------------ headings
    def h(self, text, level=1, numbered=True):
        if numbered:
            self.hnum[level - 1] += 1
            for j in range(level, 3):
                self.hnum[j] = 0
            num = ".".join(str(n) for n in self.hnum[:level]) + " "
        else:
            num = ""
        p = self.d.add_paragraph()
        p.paragraph_format.space_before = Pt(12); p.paragraph_format.keep_with_next = True
        label = num + (text.upper() if level == 1 else text)
        r = p.add_run(label); r.bold = level < 3; r.italic = level == 3
        return p

    # ------------------------------------------------------------------ equations
    def eq(self, latex):
        """Display equation as a normal paragraph: centred OMML equation, number at the right margin."""
        self.eq_n += 1
        sec = self.d.sections[0]
        width = (sec.page_width - sec.left_margin - sec.right_margin)
        p = self.d.add_paragraph()
        ts = p.paragraph_format.tab_stops
        ts.add_tab_stop(int(width / 2), WD_TAB_ALIGNMENT.CENTER)
        ts.add_tab_stop(int(width), WD_TAB_ALIGNMENT.RIGHT)
        p.add_run("\t")
        m = omml(latex)
        # display-style operators: limits of every n-ary operator (sums, integrals) above and below
        M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
        for nary in m.iter(f"{{{M}}}nary"):
            pr = nary.find(f"{{{M}}}naryPr")
            if pr is None:
                pr = OxmlElement("m:naryPr"); nary.insert(0, pr)
            loc = pr.find(f"{{{M}}}limLoc")
            if loc is None:
                loc = OxmlElement("m:limLoc"); pr.insert(0, loc)
            loc.set(f"{{{M}}}val", "undOvr")
        p._p.append(m)
        p.add_run(f"\t({self.eq_n})")
        return self.eq_n

    # ------------------------------------------------------------------ captions
    def figure(self, title, legend, file=None):
        self.fig_n += 1
        if self.embed and file is not None:
            self.picture(file, 16.5)
        p = self.d.add_paragraph()
        r = p.add_run(f"Figure {self.fig_n}. "); r.bold = True
        self.runs(p, title.rstrip(".") + ". " + legend)
        return self.fig_n

    def table(self, title, header, rows, widths_cm, legend=None, size=10, label=None):
        if label is None:
            self.tab_n += 1; label = f"Table {self.tab_n}."
        return super().table(title, header, rows, widths_cm, legend=legend, size=size, label=label)

    # ------------------------------------------------------------------ references
    def _sortkey(self, k):
        """GJI order: first author; single-author before two-author (by co-author) before et al.; then year."""
        lab, full = self.refs[k]
        year = lab.rsplit(" ", 1)[1]
        name = lab.rsplit(" ", 1)[0]
        if " et al." in name:
            return (name.replace(" et al.", "").lower(), 2, "", year, full.lower())
        if " & " in name:
            first, second = name.split(" & ")
            return (first.lower(), 1, second.lower(), year, full.lower())
        return (name.lower(), 0, "", year, full.lower())

    def references(self):
        self.h("References", 1, numbered=False)
        for k in sorted(self.used, key=self._sortkey):
            p = self.d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Cm(0.6); p.paragraph_format.first_line_indent = Cm(-0.6)
            self.runs(p, self.refs[k][1])
        return [k for k in self.refs if k not in self.used]
